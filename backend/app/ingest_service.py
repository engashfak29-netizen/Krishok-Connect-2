from __future__ import annotations
import os, re, json, csv, hashlib, mimetypes, subprocess, tempfile
from pathlib import Path
from typing import Any

BANGLA_RE = r'\u0980-\u09FF'
TOKEN_RE = re.compile(r'[\w'+BANGLA_RE+r']+', re.UNICODE)
SUPPORTED = {
    '.pdf':'pdf', '.txt':'text', '.md':'text', '.csv':'csv', '.json':'json',
    '.xlsx':'excel', '.xlsm':'excel', '.docx':'docx',
    '.jpg':'image', '.jpeg':'image', '.png':'image', '.webp':'image',
    '.mp4':'video', '.webm':'video', '.mov':'video', '.m4v':'video',
    '.mp3':'audio', '.wav':'audio', '.m4a':'audio', '.ogg':'audio'
}


def sha256_file(path: str) -> str:
    h=hashlib.sha256()
    with open(path,'rb') as f:
        for b in iter(lambda:f.read(1024*1024),b''): h.update(b)
    return h.hexdigest()


def keywords(text: str, limit=120):
    toks=TOKEN_RE.findall((text or '').lower())
    # Keep meaningful repeated tokens out, while preserving Bengali words.
    out=[]; seen=set()
    for t in toks:
        if len(t)<2 or t in seen: continue
        seen.add(t); out.append(t)
        if len(out)>=limit: break
    return out


def chunk_text(text: str, size=1400, overlap=180):
    text=re.sub(r'\r\n?', '\n', text or '').strip()
    if not text: return []
    chunks=[]; start=0
    while start < len(text):
        end=min(len(text), start+size)
        if end < len(text):
            cut=max(text.rfind('\n',start,end), text.rfind('।',start,end), text.rfind('. ',start,end))
            if cut>start+500: end=cut+1
        chunks.append(text[start:end].strip())
        if end>=len(text): break
        start=max(0,end-overlap)
    return [x for x in chunks if x]


def extract_pdf(path):
    from pypdf import PdfReader
    r=PdfReader(path); parts=[]
    for i,p in enumerate(r.pages,1):
        txt=p.extract_text() or ''
        if txt.strip(): parts.append({'text':txt,'locator':{'page':i}})
    return parts


def extract_docx(path):
    from docx import Document
    d=Document(path); text='\n'.join(p.text for p in d.paragraphs if p.text.strip())
    for t in d.tables:
        rows=[]
        for row in t.rows: rows.append(' | '.join(c.text.strip() for c in row.cells))
        if rows: text += '\n'+'\n'.join(rows)
    return [{'text':text,'locator':{}}] if text.strip() else []


def extract_excel(path):
    from openpyxl import load_workbook
    wb=load_workbook(path, read_only=True, data_only=True)
    out=[]
    for ws in wb.worksheets:
        lines=[]
        for row in ws.iter_rows(values_only=True):
            vals=[str(v).strip() for v in row if v is not None and str(v).strip()]
            if vals: lines.append(' | '.join(vals))
        if lines: out.append({'text':f'Sheet: {ws.title}\n'+'\n'.join(lines),'locator':{'sheet':ws.title}})
    return out


def extract_structured(path, kind):
    raw=Path(path).read_bytes()
    text=raw.decode('utf-8-sig',errors='replace')
    if kind=='json':
        try: text=json.dumps(json.loads(text),ensure_ascii=False,indent=2)
        except Exception: pass
    elif kind=='csv':
        rows=[]
        try:
            for row in csv.reader(text.splitlines()): rows.append(' | '.join(x.strip() for x in row))
            text='\n'.join(rows)
        except Exception: pass
    return [{'text':text,'locator':{}}] if text.strip() else []


def extract_image(path):
    # OCR is optional at runtime. If unavailable, retain a searchable metadata item and flag it.
    try:
        from PIL import Image
        import pytesseract
        lang=os.getenv('OCR_LANG','eng+ben')
        txt=pytesseract.image_to_string(Image.open(path),lang=lang)
        if txt.strip(): return [{'text':txt,'locator':{},'meta':{'ocr':True}}]
        return [{'text':'','locator':{},'meta':{'ocr':True,'empty':True}}]
    except Exception as e:
        return [{'text':'','locator':{},'meta':{'ocr':False,'ocr_error':type(e).__name__}}]


def transcribe_audio(path):
    # Optional high-quality local transcription using faster-whisper if installed.
    try:
        from faster_whisper import WhisperModel
        model_name=os.getenv('WHISPER_MODEL','small')
        model=WhisperModel(model_name, device=os.getenv('WHISPER_DEVICE','cpu'), compute_type=os.getenv('WHISPER_COMPUTE_TYPE','int8'))
        segments,_=model.transcribe(path, language=os.getenv('WHISPER_LANGUAGE','bn'), vad_filter=True)
        out=[]
        for seg in segments:
            txt=(seg.text or '').strip()
            if txt: out.append({'text':txt,'locator':{'start':seg.start,'end':seg.end},'meta':{'transcription':True}})
        return out, None
    except Exception as e:
        return [], f'faster-whisper unavailable or failed: {type(e).__name__}'


def extract_video(path):
    # Prefer direct transcription when the optional engine can read the video.
    parts,err=transcribe_audio(path)
    if parts: return parts
    return [{'text':'','locator':{},'meta':{'transcription':False,'needs_transcription':True,'transcription_error':err}}]


def extract_file(path: str):
    ext=Path(path).suffix.lower(); kind=SUPPORTED.get(ext)
    if not kind: raise ValueError(f'Unsupported file type: {ext}')
    if kind=='pdf': return extract_pdf(path), kind
    if kind=='docx': return extract_docx(path), kind
    if kind=='excel': return extract_excel(path), kind
    if kind in ('text','csv','json'): return extract_structured(path,kind), kind
    if kind=='image': return extract_image(path), kind
    if kind in ('video','audio'):
        return (extract_video(path) if kind=='video' else transcribe_audio(path)[0]), kind
    raise ValueError(kind)
